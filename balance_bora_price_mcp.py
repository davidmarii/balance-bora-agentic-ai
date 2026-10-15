
import os
from pathlib import Path
from datetime import date, datetime

from openpyxl import load_workbook
from mcp.server.mcpserver import MCPServer

mcp = MCPServer("Balance Bora Price Server")

PRICE_FILE = Path(
    os.environ.get(
        "BALANCE_BORA_PRICE_FILE",
        Path(__file__).with_name("Balance_Bora_Feed_Price_Template.xlsx")
    )
)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        try:
            return datetime.fromisoformat(str(value)).date()
        except ValueError:
            pass
    return None


@mcp.tool()
def lookup_feed_prices(
    ingredient_name: str = "",
    location: str = "",
    verified_only: bool = True
) -> dict:
    """Find dated feed ingredient prices in the Balance Bora Excel workbook.

    By default, return only prices marked 'Verified with supplier'.
    This tool does not invent prices.
    """
    if not PRICE_FILE.exists():
        return {
            "status": "file_missing",
            "message": f"Price workbook not found: {PRICE_FILE}",
            "prices": []
        }

    workbook = load_workbook(
        PRICE_FILE, data_only=True, read_only=True
    )

    if "Price List" not in workbook.sheetnames:
        workbook.close()
        return {
            "status": "sheet_missing",
            "prices": []
        }

    sheet = workbook["Price List"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)

    if not headers:
        workbook.close()
        return {"status": "empty_sheet", "prices": []}

    columns = {
        str(name).strip(): i
        for i, name in enumerate(headers)
        if name is not None
    }

    def field(row, name):
        i = columns.get(name)
        return row[i] if i is not None and i < len(row) else None

    results = []
    today = date.today()

    for row in rows:
        name = str(field(row, "Ingredient Name") or "").strip()
        supplier = str(field(row, "Supplier") or "").strip()
        county = str(field(row, "County") or "").strip()
        town = str(field(row, "Town/Market") or "").strip()
        status = str(field(row, "Verification Status") or "").strip()
        source_type = str(field(row, "Source Type") or "").strip()
        notes = str(field(row, "Notes") or "").strip()

        # Never treat test records as genuine supplier prices.
        if source_type.casefold() == "test data":
            continue
        if "test only" in notes.casefold():
            continue
        if supplier.casefold().startswith("test "):
            continue
        if not name or not supplier:
            continue
        if supplier.upper().startswith("REPLACE WITH"):
            continue
        if ingredient_name and ingredient_name.casefold() not in name.casefold():
            continue
        if location and location.casefold() not in (
            county + " " + town
        ).casefold():
            continue
        if verified_only and status.casefold() != "verified with supplier":
            continue

        try:
            pack_price = float(field(row, "Price (KES)"))
            quantity = float(field(row, "Quantity per Unit (kg)"))
            delivery_value = field(row, "Delivery Cost (KES)")
            delivery = float(delivery_value or 0)
        except (TypeError, ValueError):
            continue

        if pack_price < 0 or quantity <= 0 or delivery < 0:
            continue

        quoted = parse_date(field(row, "Date Quoted"))
        valid_until = parse_date(field(row, "Valid Until"))

        results.append({
            "ingredient": name,
            "supplier": supplier,
            "county": county,
            "town": town,
            "price_kes_per_kg": round(
                (pack_price + delivery) / quantity, 2
            ),
            "pack_price_kes": pack_price,
            "pack_size_kg": quantity,
            "date_quoted": quoted.isoformat() if quoted else None,
            "valid_until": valid_until.isoformat() if valid_until else None,
            "stale_over_7_days": (
                quoted is None or (today - quoted).days > 7
            ),
            "expired": valid_until is not None and valid_until < today,
            "verification_status": status,
            "source_type": str(field(row, "Source Type") or ""),
            "evidence": str(field(row, "Evidence / Listing URL") or "")
        })

    workbook.close()

    return {
        "status": "ok" if results else "no_matching_prices",
        "count": len(results),
        "prices": results,
        "message": (
            "Confirm old or expired quotations before using them."
            if results else
            "No matching prices found. Add real quotations to Excel."
        )
    }


if __name__ == "__main__":
    mcp.run(transport="stdio")
